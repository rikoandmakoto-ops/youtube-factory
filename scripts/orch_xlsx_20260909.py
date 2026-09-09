# -*- coding: utf-8 -*-
"""2026-09-09 指揮者 — 分析シート生成 (reports/youtube_analysis_20260909.xlsx)"""
import datetime
import os
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)
OUT = 'reports/youtube_analysis_20260909.xlsx'

CH6 = ['daily-science', 'scp-lab', '2ch-matome',
       'pokemon-lab', 'yokai-watch', 'company-facts']
QC = ','.join(f"'{c}'" for c in CH6)

con = sqlite3.connect('data/analytics/analytics.db')
SNAP = con.execute('select max(date) from video_metrics').fetchone()[0]

# ---- 母集団: 最新スナップショット・公開09-05以前(評価遅延3日)・views>=200 ----
BASE = f'''select channel_id,title,views,subscribers_gained,likes,comments,
 avg_view_percentage,published_at,video_id
 from video_metrics where date='{SNAP}'
 and substr(published_at,1,10)<='2026-09-05' and views>=200
 and channel_id in ({QC})'''
rows = con.execute(BASE).fetchall()

FONT = 'Arial'
H1 = Font(name=FONT, size=14, bold=True, color='FFFFFF')
H2 = Font(name=FONT, size=10, bold=True)
BODY = Font(name=FONT, size=10)
NOTE = Font(name=FONT, size=9, italic=True, color='595959')
FILL_T = PatternFill('solid', fgColor='1F3864')
FILL_H = PatternFill('solid', fgColor='D9E2F3')
FILL_G = PatternFill('solid', fgColor='E2EFDA')
FILL_R = PatternFill('solid', fgColor='FCE4E4')
FILL_Y = PatternFill('solid', fgColor='FFF2CC')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')

wb = Workbook()


def sheet(name, title, subtitle, widths):
    ws = wb.create_sheet(name) if wb.sheetnames != ['Sheet'] else wb.active
    ws.title = name
    ws['A1'] = title
    ws['A1'].font = H1
    ws['A1'].fill = FILL_T
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(widths))
    ws['A2'] = subtitle
    ws['A2'].font = NOTE
    ws['A2'].alignment = WRAP
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(widths))
    ws.row_dimensions[2].height = 42
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def header(ws, r, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=r, column=i, value=c)
        cell.font = H2
        cell.fill = FILL_H
        cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, vertical='center',
                                   horizontal='center')
    ws.freeze_panes = ws.cell(row=r + 1, column=1)


def put(ws, r, vals, fmts=None, fill=None):
    for i, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=i, value=v)
        cell.font = BODY
        cell.border = BOX
        if fmts and fmts[i - 1]:
            cell.number_format = fmts[i - 1]
        if fill:
            cell.fill = fill
        if isinstance(v, str) and len(v) > 30:
            cell.alignment = WRAP


LAG = ('評価母集団は video_metrics の最新スナップショット（{}取得）のうち、'
       'YouTube Analytics の 2〜3 日遅延（analytics_policy.evaluation_lag.min_age_days=3）を'
       '避けるため公開が 2026-09-05 以前、かつ views>=200 の n={}。'
       '判断軸は optimization.primary_metric = 登録者/1000再生。'
       ).format(SNAP, len(rows))

# =========================== 1. サマリー ==================================
ws = sheet('サマリー', 'YouTube Factory 指揮者レポート 2026-09-09',
           LAG + ' ／ 本日はネットワーク非到達のため新規取得は行わず、'
           '09-08 23時取得のスナップショットで分析している（trigger_20260909.command 参照）。',
           [30, 13, 13, 13, 13, 15, 46])

r = 4
ws.cell(row=r, column=1, value='本日の結論').font = Font(name=FONT, size=12, bold=True)
r += 1
CONCL = [
    ('① 「秘密」は登録を止める語だった',
     'タイトルに「秘密」を含む18本は 登録/千再生 0.120（16,609再生で登録2人）。'
     '含まない178本は 0.527。4.4倍差で、出現する4ch すべてが同方向。'
     'しかも「秘密」入りは公開が 08-30 までで古い＝登録を積む時間は長かった側なので、'
     '時間バイアスは結論と逆向きに働いている。全6chで機械ゲート禁止にした。'),
    ('② 2ch-matome の「ワイ〜」定型は再生10,972に対し登録0人',
     '直近50本のうち16本（32%）がこのテンプレで、合計 10,972 再生を集めながら'
     '登録者は 0 人。評価母集団に限っても n=11・10,781再生・登録0、'
     '同ch内の非該当は 0.282。再生は取れるが1人も変換しないので機械ゲートで禁止した。'
     '当chの全タイトルで「ハワイ」等への誤爆がゼロであることを確認済み。'),
    ('③ 🚨 scp-lab の投稿枠が二重登録され1日6本作っていた',
     'autopilot.schedule.times が [(17,0),(13,0),(19,0),(17,0),(13,0),(19,0)] と'
     '完全重複。ジョブIDが枠の添字なので、同一時刻に2回発火していた。3枠へ修正。'
     'ただし稼働中プロセスには idx=3,4,5 の古いジョブが残るため**要バックエンド再起動**。'),
    ('④ 答え提示語は語ごとに2倍の差がある',
     '実は 0.767 > 実態 0.681 > なぜ 0.568 > 正体 0.510 > 本当 0.462 > 理由 0.368。'
     '現状 daily-science / yokai-watch の repair_with 先頭は「正体」（下から3番目）。'
     '語尾に置けない「実は」は機械修復には使えないため、title_style で第一候補として指示した。'),
    ('⑤ 高評価率→登録の再現は5回目',
     '高評価率4分位で 登録/千 は Q1 0.249 → Q4 0.968（3.9倍）。'
     '一方で維持率は判断軸にしない方針を今回も追認（維持率上位の pokemon-lab 49.8% は'
     '登録/千 0.275 で下から2番目、維持率最下位の scp-lab 36.8% が 0.813 で1位）。'),
]
for t, b in CONCL:
    ws.cell(row=r, column=1, value=t).font = Font(name=FONT, size=10, bold=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = ws.cell(row=r, column=7, value=b)
    c.font = BODY
    c.alignment = WRAP
    ws.row_dimensions[r].height = 58
    r += 1

r += 1
ws.cell(row=r, column=1, value='主要6ch 実績（至上指標＝登録者/1000再生）').font = \
    Font(name=FONT, size=12, bold=True)
r += 1
hr = r
header(ws, r, ['チャンネル', '評価本数', '再生数', '登録者', '登録/千再生',
               '高評価率', '判定'])
r += 1
first = r
per = {}
for ch in CH6:
    sub = [x for x in rows if x[0] == ch]
    v = sum(x[2] for x in sub)
    s = sum(x[3] for x in sub)
    lk = sum(x[4] for x in sub)
    per[ch] = (len(sub), v, s, lk)
order = sorted(CH6, key=lambda c: -(per[c][2] / per[c][1]))
for ch in order:
    n, v, s, lk = per[ch]
    put(ws, r, [ch, n, v, s,
                f'=IF(C{r}=0,0,D{r}/C{r}*1000)',
                f'=IF(C{r}=0,0,{lk}/C{r})',
                ''],
        [None, '#,##0', '#,##0', '#,##0', '0.000', '0.00%', None])
    r += 1
last = r - 1
put(ws, r, ['合計 / 加重平均',
            f'=SUM(B{first}:B{last})', f'=SUM(C{first}:C{last})',
            f'=SUM(D{first}:D{last})',
            f'=IF(C{r}=0,0,D{r}/C{r}*1000)',
            f'=IF(C{r}=0,0,{sum(per[c][3] for c in CH6)}/C{r})', ''],
    [None, '#,##0', '#,##0', '#,##0', '0.000', '0.00%', None], FILL_H)
for i in range(1, 8):
    ws.cell(row=r, column=i).font = H2
tot = r
r += 2
ws.cell(row=r, column=1,
        value='判定列は空欄。上の登録/千再生は式で算出しているので、'
              '取得データを差し替えれば再計算される。').font = NOTE
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)

# =================== 2. チャンネル別詳細 ==================================
ws = sheet('チャンネル別詳細',
           'チャンネル別詳細 2026-09-09',
           LAG + ' CTR は video_reach_daily（09-01以降）の impressions/clicks 実測。'
           'video_metrics.ctr 列は欠損が多く分母が揃わないため使っていない。',
           [16, 8, 10, 8, 11, 10, 9, 9, 11, 40])
ctr = dict(con.execute('''select channel_id,
   round(100.0*sum(clicks)/nullif(sum(impressions),0),2)
   from video_reach_daily where date>='2026-09-01' group by channel_id''').fetchall())
imps = dict(con.execute('''select channel_id, sum(impressions)
   from video_reach_daily where date>='2026-09-01' group by channel_id''').fetchall())
r = 4
header(ws, r, ['チャンネル', '本数', '再生数', '登録者', '登録/千再生',
               '高評価率', '維持率%', 'CTR%', '表示回数', '本日の変更'])
r += 1
CHG = {
    'scp-lab': '🚨枠の重複を6→3へ / 17:00→09:00 / 「— …【ラベル】」型を禁止 / 秘密・99%を禁止',
    'company-facts': '秘密・99%を禁止 / title_style に語別ランキング（枠は09-08最適化済みで据え置き）',
    'daily-science': '秘密・99%を禁止 / viral_hooks から「99%が知らない」を削除（推奨と禁止が矛盾していた）',
    'yokai-watch': '秘密・99%を禁止 / title_style に語別ランキング（12時1.06が最良で枠は据え置き）',
    'pokemon-lab': '19:00→15:00 / theme_seeds 23→33本へ補充 / 秘密・99%を禁止',
    '2ch-matome': '「ワイ」「質問ある」定型を禁止（再生10,781・登録0）/ 07:00→09:00 / 秘密・99%を禁止',
}
for ch in order:
    n, v, s, lk = per[ch]
    sub = [x for x in rows if x[0] == ch]
    avp = sum(x[6] for x in sub) / len(sub)
    put(ws, r, [ch, n, v, s, f'=IF(C{r}=0,0,D{r}/C{r}*1000)',
                f'=IF(C{r}=0,0,{lk}/C{r})', round(avp, 1),
                ctr.get(ch, 0), imps.get(ch, 0), CHG[ch]],
        [None, '#,##0', '#,##0', '#,##0', '0.000', '0.00%', '0.0', '0.00',
         '#,##0', None])
    ws.row_dimensions[r].height = 32
    r += 1

r += 2
ws.cell(row=r, column=1, value='投稿時間帯 × 登録/千再生（ch内対照・n>=2の枠のみ）').font = \
    Font(name=FONT, size=12, bold=True)
r += 1
header(ws, r, ['チャンネル', '公開時刻(JST)', '本数', '再生数', '登録者',
               '登録/千再生', '', '', '', '現行枠か'])
r += 1


def jst(p):
    return (datetime.datetime.strptime(p[:19], '%Y-%m-%dT%H:%M:%S')
            + datetime.timedelta(hours=9)).hour


import json  # noqa: E402
cfgtimes = {}
for ch in CH6:
    d = json.load(open(f'data/channels/{ch}.json', encoding='utf-8'))
    cfgtimes[ch] = {t['hour'] for t in d['autopilot']['schedule']['times']}
for ch in order:
    b = {}
    for x in [y for y in rows if y[0] == ch]:
        h = jst(x[7])
        b.setdefault(h, [0, 0, 0])
        b[h][0] += 1
        b[h][1] += x[2]
        b[h][2] += x[3]
    for h, (n, v, s) in sorted(b.items(), key=lambda kv: -(kv[1][2] / kv[1][1])):
        if n < 2:
            continue
        cur = '★現行枠' if h in cfgtimes[ch] else ''
        f = FILL_G if cur and (1000 * s / v) >= 0.5 else (
            FILL_R if cur and (1000 * s / v) < 0.25 else None)
        put(ws, r, [ch, f'{h}時', n, v, s, f'=IF(D{r}=0,0,E{r}/D{r}*1000)',
                    '', '', '', cur],
            [None, None, '#,##0', '#,##0', '#,##0', '0.000',
             None, None, None, None], f)
        r += 1

# =================== 3. 動画別パフォーマンス ==============================
ws = sheet('動画別パフォーマンス', '動画別パフォーマンス 2026-09-09',
           LAG + ' 登録/千再生の降順。「本日の禁止ゲートに該当」列は、'
           '今日追加した機械ゲート（秘密／99%／ワイ・質問ある／記録票フォーマット）に'
           '引っかかるタイトルを示す。',
           [15, 62, 9, 9, 8, 8, 12, 9, 13, 22])
r = 4
header(ws, r, ['チャンネル', 'タイトル', '再生数', '登録者', '高評価', 'コメント',
               '登録/千再生', '維持率%', '公開日時(JST)', '本日の禁止ゲートに該当'])
r += 1
import re  # noqa: E402


def gate(ch, t):
    hit = []
    if '秘密' in t:
        hit.append('秘密')
    if re.search(r'(?:9\s*9|９\s*９)\s*[%％]', t):
        hit.append('99%型')
    if ch == '2ch-matome' and (re.search(r'(?<![ハスロ])ワイ(?![ドヤンルフパブザ])', t)
                               or '質問ある' in t):
        hit.append('ワイ/質問ある')
    if ch == 'scp-lab' and re.search(r'—\s*.*【', t):
        hit.append('記録票フォーマット')
    return ' / '.join(hit)


for x in sorted(rows, key=lambda y: -(1000 * y[3] / y[2])):
    ch, t, v, s, lk, cm, avp, pub, _vid = x
    g = gate(ch, t)
    pj = (datetime.datetime.strptime(pub[:19], '%Y-%m-%dT%H:%M:%S')
          + datetime.timedelta(hours=9)).strftime('%m-%d %H:%M')
    put(ws, r, [ch, t, v, s, lk, cm, f'=IF(C{r}=0,0,D{r}/C{r}*1000)',
                round(avp, 1), pj, g],
        [None, None, '#,##0', '#,##0', '#,##0', '#,##0', '0.000', '0.0',
         None, None],
        FILL_R if g else None)
    r += 1

# =================== 4. 改善アクション ====================================
ws = sheet('改善アクション', '改善アクション 2026-09-09',
           'すべて実データ根拠つき。「反証条件」を満たしたら次の指揮者が撤回すること。'
           'コンフィグへの反映は完了済み（data/channels/*.json。'
           'channels_orchestrator/ は 09-08 から symlink なので書き込み先は1箇所）。',
           [4, 15, 34, 40, 34, 14, 12])
r = 4
header(ws, r, ['#', '対象', '施策', '実データ根拠', '反証条件', '評価予定日', '状態'])
r += 1
ACTIONS = [
    ('全6ch', 'title_rules.hard_constraints.banned_words に「秘密」を追加',
     '登録/千 0.120 (n=18/16,609再生/登録2) vs 0.527 (n=178)。'
     'pokemon 0.122vs0.332・yokai 0.149vs0.543・daily 0.00vs0.497・2ch 0.00vs0.209 と'
     '出現4ch全て同方向。「秘密」入りは公開08-30までで古い＝時間バイアスは逆向き。',
     '09-16 時点で「秘密」を含まない新規コホートの 登録/千 が 0.35 を下回ったら、'
     '語ではなく別要因だったとして撤回',
     '2026-09-16', '反映済'),
    ('全6ch', 'forbid_patterns に 99%型（(?:9\\s*9|９\\s*９)\\s*[%％]）を追加',
     '登録/千 0.194 (n=11) vs 0.510 (n=185)。',
     'n=11 と小さい。09-20 までに該当0本のまま推移し全体の登録/千が改善しなければ'
     '効果無しとみなす', '2026-09-20', '反映済'),
    ('2ch-matome', 'forbid_patterns に「ワイ」（位置不問）「質問ある」を追加',
     '該当 n=11 で合計 10,781 再生・登録者 0 人。同ch内の非該当は 0.282 (n=27)。'
     '直近50本では16本(32%)が該当し10,972再生・登録0。再生は取れるが変換が完全にゼロ。'
     '当初は行頭限定にしたが n=7/11 しか止まらないため位置不問へ拡張し、'
     '「ハワイ」等への誤爆が当ch全タイトルでゼロであることを確認した。',
     '禁止後に当chの総再生が 3割以上落ち、かつ登録/千が 0.28 を超えなければ、'
     '再生の犠牲に見合わないとして撤回', '2026-09-16', '反映済'),
    ('scp-lab', '🚨 autopilot.schedule.times の重複を 6枠→3枠へ修正',
     'times が [(17,0),(13,0),(19,0),(17,0),(13,0),(19,0)]。'
     '_job_id(channel_id, idx) が枠の添字なので同一時刻に2ジョブが登録され、'
     '1日6本作っていた。',
     '（バグ修正のため反証条件なし）',
     '—', '要バックエンド再起動'),
    ('scp-lab', '投稿枠 17:00 → 09:00',
     '当ch内 19時 1.23(n=9) / 13時 0.83(n=6) / 9時 0.69(n=12)。'
     '17時は n<2 で有効データが無かった。',
     '09-20 時点で 9時枠の 登録/千 が 0.4 を下回ったら 12時(0.80,n=2)へ',
     '2026-09-20', '反映済'),
    ('scp-lab', 'forbid_patterns に「— …【ラベル】」記録票フォーマットを追加',
     '当ch内で該当 n=6・5,418再生・登録者0人、非該当 n=30 で 0.960。'
     '一方 SCP-数字ID を含む型は 0.888(n=27) vs 含まない 0.449(n=9)。',
     'n=6 と小さい。09-20 までに非該当コホートの 登録/千 が 0.5 を下回ったら撤回',
     '2026-09-20', '反映済'),
    ('2ch-matome', '投稿枠 07:00 → 09:00',
     '当ch内 7時 0.00(n=2)が最下位、9時 0.29(n=5)・12時 0.29(n=7)が最上位。',
     '7時は n=2 と弱い。09-20 に 9時枠が 0.2 を下回ったら 17時(0.28,n=3)へ',
     '2026-09-20', '反映済'),
    ('pokemon-lab', '投稿枠 19:00 → 15:00',
     '当ch内 19時は n<2 でデータ無し。隣接の18時は 0.21(n=10) で最下位。'
     '15時 0.33(n=3) / 17時 0.31(n=10) が上位。8時 0.84 は n=2 で採用せず。',
     '15時は n=3 と弱い。09-20 に 0.25 を下回ったら 8時(0.84,n=2)を試す',
     '2026-09-20', '反映済'),
    ('pokemon-lab', 'theme_seeds を 23 → 33 本へ補充',
     'theme_queue が 12本と6ch中最少（他は22〜46本）。'
     'queue は seeds から30分ごとに再生成される揮発キャッシュなので seeds 側に投入。'
     '題材は上位実績の「固有名詞＋数値を1つ提示」型に寄せ、'
     '0転換だった「どっちが勝つ？」対戦比較型は増やしていない。',
     '補充分が公開されて 登録/千 0.28（当ch現状）を下回ったら題材選定を見直す',
     '2026-09-20', '反映済'),
    ('全6ch', 'theme_priority.title_style に答え提示語の語別ランキングを反映',
     '実は 0.767(n=18) > 実態 0.681(n=9) > なぜ 0.568(n=72) > 正体 0.510(n=25) > '
     '本当 0.462(n=17) > 理由 0.368(n=19)。'
     '答え提示語あり 0.606(n=123) vs なし 0.302(n=73)。',
     '自然文の title_style は LLM に無視されることが 09-04 に実測確定している。'
     '09-16 に「実は」の出現率が上がっていなければ、hard_constraints 側の'
     'repair_with 並べ替えへ切り替える', '2026-09-16', '反映済'),
    ('daily-science', 'theme_priority.viral_hooks から「99%が知らない」系を削除',
     '同ch設定が、実測で劣位（0.194 vs 0.510）かつ本日禁止した語を'
     '「バイラルフック」として推奨していた。推奨と機械ゲートの矛盾を解消。',
     '（矛盾解消のため反証条件なし）', '—', '反映済'),
    ('backend', '【コード変更提案・未実施】title_constraints._REASON_WORDS から「秘密」を外す',
     '同モジュールの _REASON_WORDS に「秘密」が含まれており、'
     '「なぜ」始まりの機械的な言い換えの際に「秘密」があると理由語ありと判定して'
     '補わない。本日の実測では「秘密」は最悪の語なので、この判定は逆効果。',
     '夜間・テスト無しでのコード変更は避けたため未実施。'
     'banned_words による停止が先に効くので実害は限定的。',
     '—', '未実施（要人手）'),
    ('backend', '【コード変更提案・未実施】hard_constraints に require_pattern を追加',
     'scp-lab は SCP-数字ID を含む型が 0.888(n=27) vs 含まない 0.449(n=9) で2倍。'
     'しかし require_any_of は語彙照合のみで正規表現の必須化を表現できない。',
     '現状は theme_priority.title_style の自然文で指示しているが、'
     '自然文は無視されることが実測済みなので、機械ゲート化が要る。',
     '—', '未実施（要人手）'),
]
for i, a in enumerate(ACTIONS, 1):
    fill = FILL_Y if a[5].startswith('未実施') or '再起動' in a[5] else None
    put(ws, r, [i] + list(a), [None] * 7, fill)
    ws.row_dimensions[r].height = 60
    r += 1

# =========================== 5. トレンド ==================================
ws = sheet('トレンド', 'トレンド 2026-09-09',
           LAG + ' 各切り口は同一母集団(n=196)の分割。'
           'n が小さい行は参考値として扱うこと。',
           [34, 30, 9, 11, 9, 13, 46])
r = 4


def block(title, note, items):
    global r
    ws.cell(row=r, column=1, value=title).font = Font(name=FONT, size=12, bold=True)
    r += 1
    if note:
        c = ws.cell(row=r, column=1, value=note)
        c.font = NOTE
        c.alignment = WRAP
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        ws.row_dimensions[r].height = 28
        r += 1
    header(ws, r, ['切り口', '区分', '本数', '再生数', '登録者', '登録/千再生', '読み方'])
    r += 1
    for label, seg, read in items:
        v = sum(x[2] for x in seg)
        s = sum(x[3] for x in seg)
        put(ws, r, [title, label, len(seg), v, s,
                    f'=IF(D{r}=0,0,E{r}/D{r}*1000)', read],
            [None, None, '#,##0', '#,##0', '#,##0', '0.000', None])
        ws.row_dimensions[r].height = 26
        r += 1
    r += 1


MK = ['理由', '正体', '本当の', '実は', 'わけ', 'なぜ', '真相', '裏側', '実態']
rs = sorted(rows, key=lambda x: x[4] / x[2])
q = len(rs) // 4
block('高評価率4分位', '5回目の再現。維持率とは違い、登録転換と単調に近い関係がある。',
      [(f'Q{i+1}', rs[i * q:(i + 1) * q] if i < 3 else rs[3 * q:],
        'Q4はQ1の3.9倍' if i == 3 else '') for i in range(4)])

block('答え提示語（語別）', '語尾に置けない「実は」は repair_with に入れられない。'
      'title_style 側で第一候補として指示している。',
      [(m, [x for x in rows if m in x[1]],
        '最上位' if m == '実は' else ('最下位' if m == '理由' else ''))
       for m in ['実は', '実態', 'なぜ', '正体', '本当', '理由']
       if len([x for x in rows if m in x[1]]) >= 8])

block('本日禁止した型', '3つとも本日 hard_constraints へ反映済み。',
      [('「秘密」あり', [x for x in rows if '秘密' in x[1]], '4.4倍の劣位'),
       ('「秘密」なし', [x for x in rows if '秘密' not in x[1]], ''),
       ('99%型', [x for x in rows if re.search(r'99\s*[%％]', x[1])], '2.6倍の劣位'),
       ('99%型以外', [x for x in rows if not re.search(r'99\s*[%％]', x[1])], ''),
       ('2ch「ワイ/質問ある」',
        [x for x in rows if x[0] == '2ch-matome'
         and (re.search(r'(?<![ハスロ])ワイ(?![ドヤンルフパブザ])', x[1])
              or '質問ある' in x[1])], '登録0人'),
       ('2ch その他',
        [x for x in rows if x[0] == '2ch-matome'
         and not (re.search(r'(?<![ハスロ])ワイ(?![ドヤンルフパブザ])', x[1])
                  or '質問ある' in x[1])], '')])

block('scp-lab のタイトル型', '固有IDの提示が効く一方、記録票フォーマットは劣位。',
      [('SCP-数字ID あり',
        [x for x in rows if x[0] == 'scp-lab' and re.search(r'SCP-\d+', x[1])], '2.0倍'),
       ('SCP-数字ID なし',
        [x for x in rows if x[0] == 'scp-lab' and not re.search(r'SCP-\d+', x[1])], ''),
       ('「— …【ラベル】」型',
        [x for x in rows if x[0] == 'scp-lab' and re.search(r'—\s*.*【', x[1])], '本日禁止'),
       ('それ以外',
        [x for x in rows if x[0] == 'scp-lab' and not re.search(r'—\s*.*【', x[1])], '')])

block('維持率と登録転換', '維持率を判断軸にしない方針の再確認（09-04決定・09-08追認）。',
      [('維持率50%以上', [x for x in rows if x[6] >= 50], ''),
       ('維持率50%未満', [x for x in rows if x[6] < 50], '差はほぼ無い')])

waste = [x for x in rows if x[2] >= 1200 and x[3] == 0]
block('高再生・ゼロ転換', f'再生は集めているのに登録が1人も出ていない層。'
      f'{len(waste)}本で合計 {sum(x[2] for x in waste):,} 再生。'
      f'本日の禁止ゲートはこの層を狙って設計している。',
      [('views>=1200 かつ 登録0', waste, '最大の取りこぼし'),
       ('それ以外', [x for x in rows if not (x[2] >= 1200 and x[3] == 0)], '')])

wb.save(OUT)
print('saved:', OUT)
