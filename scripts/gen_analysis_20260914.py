# -*- coding: utf-8 -*-
"""2026-09-14 指揮者タスク: analytics.db から日次分析 xlsx を生成する。

データ源は data/analytics/analytics.db のみ。ブラウザでの YouTube 直接アクセス・
API キーは使わない。
"""
import sqlite3, datetime, collections, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB = os.path.join(ROOT, 'data/analytics/analytics.db')
OUT = os.path.join(ROOT, 'reports/youtube-analysis-2026-09-14.xlsx')
SNAP, PREV = '2026-09-13', '2026-09-08'

c = sqlite3.connect(DB); c.row_factory = sqlite3.Row
rows = [dict(r) for r in c.execute("select * from video_metrics where date=?", (SNAP,))]
prow = [dict(r) for r in c.execute("select * from video_metrics where date=?", (PREV,))]

NAMES = {'scp-lab': 'ゆっくり異常存在SCPラボ', 'daily-science': 'リコとマコトのゆっくり日常科学',
         'pokemon-lab': 'ゆっくりポケラボ', 'yokai-watch': 'ゆっくり妖怪ラボ',
         '2ch-matome': 'ゆっくり2chスレまとめ劇場', 'company-facts': '企業のホンネ',
         'clip-lab': '切り抜きラボ（ひろゆき）', 'clip-fukada': '深田えいみ 切り抜き',
         'clip-kaneko': '金子みゆ 切り抜き', 'fake-paper': '虚構論文チャンネル',
         'akashic-librarian': 'ラグナロクの司書', 'clip-animal': '動物情報局',
         'socio-rx': '社会学の処方箋'}
ORDER = ['scp-lab', 'yokai-watch', 'company-facts', 'daily-science', 'socio-rx', '2ch-matome',
         'pokemon-lab', 'akashic-librarian', 'clip-fukada', 'clip-kaneko', 'fake-paper',
         'clip-lab', 'clip-animal']
ACTIVE = {'scp-lab', 'yokai-watch', 'company-facts', 'daily-science', 'socio-rx'}
THUMB_OK = {'daily-science'}
THUMB_NG = {'scp-lab', 'company-facts', 'pokemon-lab', 'yokai-watch', 'fake-paper'}
CLIPS = ('clip-lab', 'clip-fukada', 'clip-kaneko', 'clip-animal')


def agg(rs, ch):
    s = [r for r in rs if r['channel_id'] == ch]
    v = sum(r['views'] or 0 for r in s); l = sum(r['likes'] or 0 for r in s)
    sg = sum(r['subscribers_gained'] or 0 for r in s); cm = sum(r['comments'] or 0 for r in s)
    ret = [r['avg_view_percentage'] for r in s if r['avg_view_percentage']]
    return dict(n=len(s), v=v, l=l, sg=sg, cm=cm,
                spm=1000 * sg / v if v else 0, lr=100 * l / v if v else 0,
                ret=sum(ret) / len(ret) if ret else 0)


def latest_agg(ch):
    d = c.execute("select max(date) from video_metrics where channel_id=?", (ch,)).fetchone()[0]
    if not d:
        return None, None
    rs = [dict(r) for r in c.execute(
        "select * from video_metrics where channel_id=? and date=?", (ch, d))]
    return d, agg(rs, ch)


H1 = Font(bold=True, size=15); H2 = Font(bold=True, size=12, color='FFFFFF')
HDR = PatternFill('solid', fgColor='1F3864'); SUB = PatternFill('solid', fgColor='D9E2F3')
RED = PatternFill('solid', fgColor='FCE4E4'); GRN = PatternFill('solid', fgColor='E2EFDA')
YEL = PatternFill('solid', fgColor='FFF2CC'); GRY = PatternFill('solid', fgColor='F2F2F2')
B = Border(*[Side('thin', color='BFBFBF')] * 4)


def hrow(ws, r, vals):
    for i, v in enumerate(vals, 1):
        cl = ws.cell(r, i, v); cl.font = H2; cl.fill = HDR; cl.border = B
        cl.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[r].height = 28


def setw(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = openpyxl.Workbook()

# ============================== サマリ ==============================
ws = wb.active; ws.title = 'サマリ'
setw(ws, [30, 16, 16, 14, 14, 58])
ws['A1'] = 'YouTube Factory 日次分析  2026-09-14（月）'; ws['A1'].font = H1
ws['A2'] = 'データ源: data/analytics/analytics.db（最終 fetch 2026-09-13 22:40 JST）／ブラウザ・APIキーは不使用'
ws['A2'].font = Font(italic=True, size=9, color='666666')

r = 4
ws.cell(r, 1, '★ 本日の結論').font = Font(bold=True, size=13); r += 1
for line, fill in [
    ('① 公開は完全に再開した。09-13に7本、09-14未明に4本、計11本。5日間の停止は解消済み。', GRN),
    ('② 【新規P0】サムネイルが5chで1枚も適用されていない。403 "authenticated user doesn\'t have '
     'permissions to upload and set custom video thumbnails"。全ログで失敗104件／成功は '
     'daily-science の16件のみ。', RED),
    ('③ 09-13公開の7本は analytics 上まだ全て views=0（YouTube側の集計ラグ）。'
     'よって「新しい動画の実績」はまだ1本も測れていない。', YEL),
    ('④ config は本日バックエンドのオーケストレータが既に更新済み（投稿時間・型優先度／5ch）。'
     '同一スナップショットでの二重判断を避けるため、指揮者からの追加変更は行わず、'
     '内容を独立に検算したうえでコミットした。', GRN),
    ('⑤ 【新規バグ修正】title_constraints.repair() が forbid_patterns を一切見ておらず、'
     '規約違反タイトルがそのまま公開されていた（実例: 09-13 23:15 company-facts「…FC比率99%…」）。'
     '本日修正し、テストの無退行を確認。', GRN),
]:
    cl = ws.cell(r, 1, line); cl.fill = fill
    cl.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 34; r += 1

r += 1
ws.cell(r, 1, '稼働5chの主要KPI（09-13スナップショット・直近50本ローリング窓）').font = Font(bold=True, size=12)
r += 1
hrow(ws, r, ['チャンネル', '本数', '総再生', '登録', '登録/千再生', '高評価率%']); r += 1
tv = tl = ts = 0
for ch in ['scp-lab', 'yokai-watch', 'company-facts', 'daily-science', 'socio-rx']:
    a = agg(rows, ch); tv += a['v']; tl += a['l']; ts += a['sg']
    for i, v in enumerate([NAMES[ch], a['n'], a['v'], a['sg'], round(a['spm'], 3),
                           round(a['lr'], 3)], 1):
        ws.cell(r, i, v).border = B
    r += 1
for i, v in enumerate(['合計 / 加重平均', '', tv, ts, round(1000 * ts / tv, 3),
                       round(100 * tl / tv, 3)], 1):
    cl = ws.cell(r, i, v); cl.font = Font(bold=True); cl.fill = SUB; cl.border = B
r += 2

ws.cell(r, 1, '至上指標「登録/千再生」の系統比較').font = Font(bold=True, size=12); r += 1
hrow(ws, r, ['系統', '本数', '総再生', '登録', '登録/千再生', '注記']); r += 1
yk = ['scp-lab', 'yokai-watch', 'company-facts', 'daily-science', 'socio-rx', '2ch-matome',
      'pokemon-lab']
ratios = {}
for label, chs, note in [
        ('ゆっくり系 7ch', yk, '稼働5ch＋停止2ch。各chの最新スナップショット'),
        ('切り抜き系 4ch', list(CLIPS), '全て停止中。最終データ 09-08 / 09-06')]:
    V = S = N = 0
    for ch in chs:
        _d, a = latest_agg(ch)
        if a:
            V += a['v']; S += a['sg']; N += a['n']
    ratios[label] = 1000 * S / V if V else 0
    for i, v in enumerate([label, N, V, S, round(ratios[label], 3), note], 1):
        ws.cell(r, i, v).border = B
    r += 1
mult = ratios['ゆっくり系 7ch'] / ratios['切り抜き系 4ch'] if ratios['切り抜き系 4ch'] else 0
cl = ws.cell(r, 1, f'→ 倍率 {mult:.2f}倍（ゆっくり系が優位）。'
                   '※窓とch構成で数字が動くため、倍率を書くときは必ず両方を併記すること'
                   '（09-11〜09-13で3回誤読が起きている）')
cl.font = Font(italic=True, size=9)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6); r += 2

ws.cell(r, 1, '★ ブロッカー（優先順）').font = Font(bold=True, size=13); r += 1
hrow(ws, r, ['#', '優先', '内容', '実行者', '期日', '根拠 / 補足']); r += 1
for row_ in [
    (1, 'P0', 'サムネイル権限。scp-lab / company-facts / pokemon-lab / yokai-watch / fake-paper '
              'の5chで custom thumbnail が403。各chのYouTubeアカウントを電話認証'
              '（youtube.com/verify）する必要がある', 'ユーザー操作（自動不可）', '即',
     '「サムネ品質最優先」の施策が、そもそも1枚も反映されていない。生成コストが丸ごと捨てられている'),
    (2, 'P0', 'GCP OAuth同意画面を「テスト中」→「本番」へ公開（project 844705815004）',
     'ユーザー操作（自動不可）', '09-20まで',
     'テスト中は refresh token が7日で失効。09-13再認可の6chは09-20前後に一斉に落ちる'),
    (3, 'P0', '残り7chの再認可。特に akashic-librarian（登録/千 0.571 = 上位）',
     'ユーザー操作（自動不可）', '—', '停止理由は実力ではなくOAuth'),
    (4, 'P1', 'ANTHROPIC_API_KEY が401（invalid）。dual scenario gen が GPT単独に劣化中',
     'ユーザー操作', '—',
     'backend.log: "Claude call failed … API key is invalid"。09-13夜の「有効化」が効いていない'),
    (5, 'P1', 'yokai-watch の 19:00枠。本日の自動変更で 17:00→19:00 に移設されたが、'
              '独立検算では 19時 0.41(n=21) が当ch2番目に弱い枠。'
              '12時 1.21(n=12) / 16時 1.65(n=3) のほうが強い', '指揮者', '09-15',
     'バックエンドは 19時 0.502(n=10) を根拠にした。窓の取り方で n が倍違う。'
     '同一スナップショットでの上書きは避け、次の実績で決着させる'),
    (6, 'P2', 'pokemon-lab の theme_queue に規約違反テーマが5件残存'
              '（test_enforced_channels_have_clean_queues_and_seeds が failing）', '指揮者', '再開時',
     '停止中chのため実害は先送り可。09-14の変更前後で failing 状況は同一'),
    (7, 'P2', 'logs/backend.log 84.5MB・ローテーション未実装', '—', '—',
     '09-13 84.2MB → 09-14 84.5MB'),
]:
    for i, v in enumerate(row_, 1):
        cl = ws.cell(r, i, v); cl.border = B
        cl.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(r, 2).fill = RED if row_[1] == 'P0' else (YEL if row_[1] == 'P1' else GRY)
    ws.row_dimensions[r].height = 48; r += 1
setw(ws, [6, 8, 52, 22, 12, 58])
ws.freeze_panes = 'A5'

# ========================= チャンネル別詳細 =========================
ws = wb.create_sheet('チャンネル別詳細')
setw(ws, [20, 26, 8, 12, 8, 10, 8, 12, 11, 10, 10, 10, 40])
ws['A1'] = 'チャンネル別詳細（全13ch・各chの最新スナップショット / 直近50本ローリング窓）'
ws['A1'].font = H1
ws['A2'] = ('※ CTR列は掲載しない。impressions が窓と整合せず 100% を超える値が出るため。'
            'CTRを見るときは video_reach_daily を使うこと')
ws['A2'].font = Font(italic=True, size=9, color='C00000')
r = 4
hrow(ws, r, ['channel_id', 'チャンネル名', '状態', 'データ日', '本数', '総再生', '登録',
             '登録/千再生', '高評価率%', 'コメント', '維持率%', 'サムネ', '備考']); r += 1
NOTES = {
    'socio-rx': '09-13に初公開1本。再生0のため判定不能',
    'clip-lab': '42,321再生で登録1。再生は目的関数ではない実例',
    'fake-paper': '7,047再生で登録0',
    'akashic-librarian': '停止理由はOAuthであって実力ではない',
    'clip-animal': '総再生17。実質未稼働',
    '2ch-matome': 'ゆっくり系で最下位。09-13に稼働停止',
}
for ch in ORDER:
    d, a = latest_agg(ch)
    if a is None:
        continue
    stale = (d != SNAP)
    note = []
    if stale:
        note.append(f'データが{d}で止まっている（OAuth失効）')
    if ch in NOTES:
        note.append(NOTES[ch])
    vals = [ch, NAMES[ch], '稼働' if ch in ACTIVE else '停止', d, a['n'], a['v'], a['sg'],
            round(a['spm'], 3), round(a['lr'], 3), a['cm'], round(a['ret'], 1),
            '403' if ch in THUMB_NG else ('OK' if ch in THUMB_OK else '—'), ' / '.join(note)]
    for i, v in enumerate(vals, 1):
        cl = ws.cell(r, i, v); cl.border = B
        cl.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(r, 3).fill = GRN if ch in ACTIVE else GRY
    if stale:
        ws.cell(r, 4).fill = RED
    ws.cell(r, 12).fill = RED if ch in THUMB_NG else (GRN if ch in THUMB_OK else GRY)
    if a['spm'] >= 0.6:
        ws.cell(r, 8).fill = GRN
    elif a['spm'] < 0.2:
        ws.cell(r, 8).fill = RED
    r += 1

r += 1
ws.cell(r, 1, '投稿時刻別 登録/千再生（ゆっくり系横断・09-13スナップショット・views>=200）'
        ).font = Font(bold=True, size=12); r += 1
hrow(ws, r, ['時刻(JST)', '本数', '総再生', '登録', '登録/千再生', '評価']); r += 1
ag = collections.defaultdict(lambda: [0, 0, 0])
for x in rows:
    if x['channel_id'] in CLIPS + ('fake-paper', 'akashic-librarian'):
        continue
    v = x['views'] or 0
    if v < 200 or not x['published_at']:
        continue
    t = datetime.datetime.strptime(x['published_at'][:19], '%Y-%m-%dT%H:%M:%S') \
        + datetime.timedelta(hours=9)
    a = ag[t.hour]; a[0] += v; a[1] += x['subscribers_gained'] or 0; a[2] += 1
for h in sorted(ag):
    v, s, n = ag[h]; spm = 1000 * s / v if v else 0
    ev = '◎ 最良帯' if spm >= 0.8 and n >= 10 else (
        '○' if spm >= 0.5 else ('× 弱い' if n >= 10 else 'n不足'))
    for i, x in enumerate([f'{h:02d}時', n, v, s, round(spm, 3), ev], 1):
        ws.cell(r, i, x).border = B
    if spm >= 0.8 and n >= 10:
        ws.cell(r, 5).fill = GRN
    elif n >= 10 and spm < 0.35:
        ws.cell(r, 5).fill = RED
    r += 1
ws.freeze_panes = 'A5'

# ========================== 直近動画一覧 ==========================
ws = wb.create_sheet('直近動画一覧')
setw(ws, [17, 20, 54, 10, 8, 8, 12, 11, 10])
ws['A1'] = '直近公開動画一覧（公開 2026-08-31 以降・09-13スナップショット時点の実測）'
ws['A1'].font = H1
ws['A2'] = '※ 09-13公開の7本は全て views=0。YouTube Analytics の集計ラグであり、失敗ではない'
ws['A2'].font = Font(italic=True, size=9, color='C00000')
r = 4
hrow(ws, r, ['公開日時(JST)', 'チャンネル', 'タイトル', '再生', 'いいね', '登録',
             '登録/千再生', '高評価率%', '維持率%']); r += 1
rec = sorted([x for x in rows if (x['published_at'] or '') >= '2026-08-31'],
             key=lambda x: x['published_at'], reverse=True)
for x in rec:
    t = datetime.datetime.strptime(x['published_at'][:19], '%Y-%m-%dT%H:%M:%S') \
        + datetime.timedelta(hours=9)
    v = x['views'] or 0; sg = x['subscribers_gained'] or 0; l = x['likes'] or 0
    title = (x['title'] or '').split(' #shorts')[0].strip()
    vals = [t.strftime('%m-%d %H:%M'), x['channel_id'], title, v, l, sg,
            round(1000 * sg / v, 3) if v else '—', round(100 * l / v, 3) if v else '—',
            round(x['avg_view_percentage'] or 0, 1)]
    for i, val in enumerate(vals, 1):
        ws.cell(r, i, val).border = B
    if v == 0:
        for i in range(1, 10):
            ws.cell(r, i).fill = GRY
    elif 1000 * sg / v >= 2.0:
        ws.cell(r, 7).fill = GRN
    elif v >= 1000 and sg == 0:
        ws.cell(r, 7).fill = RED
    r += 1
ws.freeze_panes = 'A5'
ws.auto_filter.ref = f'A4:I{r - 1}'

# ============================ 改善提案 ============================
ws = wb.create_sheet('改善提案')
setw(ws, [5, 10, 28, 46, 50, 22])
ws['A1'] = '改善提案 / 本日実施したアクション'; ws['A1'].font = H1
r = 3
ws.cell(r, 1, '【A】本日 実際に反映した変更（コミット済み）').font = Font(bold=True, size=12); r += 1
hrow(ws, r, ['#', '種別', '対象', '変更内容', '根拠（実測）', '検証期日']); r += 1
A = [
    (1, 'コード', 'backend/pipeline/title_constraints.py',
     'repair() に forbid_patterns の除去処理を追加。一致箇所を落として整形し、'
     '日本語が壊れる場合は既存ガードで原文に戻す',
     '09-13 23:15 company-facts「コメダ珈琲がFC比率99%にする本当の理由とは」が check で '
     'ok:false と判定されながら公開された。repair() は forbid_prefixes / 数字 / banned_words / '
     'max_chars しか見ておらず forbid_patterns を素通りさせていた。修正後は 99% と 📈 の除去を確認。'
     'テストは 110 passed / 3 failed で変更前と完全に同一（無退行）', '09-15の新規公開分'),
    (2, 'config', 'daily-science', '投稿枠 7:30/12:30/17:00 → 7:30/15:00/17:00',
     '12:30枠 0.180(n=6) に対し 17:00枠 1.236(n=14) で6.9倍。'
     '独立検算でも 12時 0.18(n=6) / 17時 1.13(n=17) を再現', '09-20'),
    (3, 'config', 'scp-lab', '投稿枠 9:00/13:00/19:00 → 13:00/17:00/19:00',
     '9:00枠 0.520(n=8) が当ch最下位。独立検算 09時 0.52(n=10) / 17時 1.32(n=4) / 19時 1.19(n=17)',
     '09-20'),
    (4, 'config', 'company-facts', '投稿枠 8:15/12:30/17:00/19:00 → 12:30/15:00/17:00/19:00',
     '8:15枠 0.235(n=4)。独立検算 08時 0.24(n=4) / 17時 0.82(n=15) / 19時 0.89(n=6)。'
     '8:15の廃止は妥当', '09-20'),
    (5, 'config', 'pokemon-lab（停止中）',
     '投稿枠 8:30/15:00/17:00 → 12:30/15:00/17:00、型を C（種族値など具体数値の提示）優先へ',
     'C型 0.53(n=9) > B型 0.35(n=6) > A型 0.17(n=16)。疑問フレームは当chでは登録に変換しない',
     '再開後'),
    (6, 'config', '稼働5ch＋pokemon-lab',
     'theme_queue を「目標型」順に再配置（scp-lab=A型 / daily-science=A型 / company-facts=C型 / '
     'pokemon-lab=C型）。新規補充は0件',
     'scp-lab: A型 1.00(n=25) vs C型 0.49(n=26) で2.0倍。'
     'company-facts: C型 0.66(n=31) ≧ A型 0.55(n=4) で現行維持', '09-20'),
]
for row_ in A:
    for i, v in enumerate(row_, 1):
        cl = ws.cell(r, i, v); cl.border = B
        cl.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(r, 2).fill = YEL if row_[1] == 'コード' else GRN
    ws.row_dimensions[r].height = 70; r += 1

r += 1
ws.cell(r, 1, '【B】あえて今日は変えなかったもの（理由つき）').font = Font(bold=True, size=12); r += 1
hrow(ws, r, ['#', '種別', '対象', '見送った変更', '理由', '再判定日']); r += 1
for row_ in [
    (1, 'config', 'yokai-watch', '17:00→19:00 の移設を差し戻すこと',
     '独立検算では 19時 0.41(n=21) が当ch2番目に弱く、12時 1.21(n=12)・16時 1.65(n=3) のほうが強い。'
     'ただしバックエンドは同じ日に 19時 0.502(n=10) で判断しており、'
     '同一スナップショット上で判断を上書きすると振動する。次の実績で決着させる', '09-15'),
    (2, 'config', '全ch', '投稿本数のさらなる増加',
     '09-12に scp-lab 週7日化・company-facts 枠4増設を入れたが、'
     '公開が止まっていたため効果が1本も測れていない。測る前に積み増さない', '09-20'),
    (3, 'config', 'socio-rx', 'テーマ優先度・投稿時間の調整', '公開1本・再生0。判定材料がゼロ',
     '初回実績が出てから'),
]:
    for i, v in enumerate(row_, 1):
        cl = ws.cell(r, i, v); cl.border = B
        cl.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(r, 2).fill = GRY; ws.row_dimensions[r].height = 62; r += 1

r += 1
ws.cell(r, 1, '【C】ユーザー操作が必要（自動化不可）').font = Font(bold=True, size=12); r += 1
hrow(ws, r, ['#', '優先', '対象', 'やること', '効かせたい指標', '放置した場合']); r += 1
for row_ in [
    (1, 'P0', 'scp-lab / company-facts / pokemon-lab / yokai-watch / fake-paper',
     '各chのYouTubeアカウントで電話認証（youtube.com/verify）を通し、カスタムサムネイル権限を得る',
     'CTR → 再生 → 登録',
     'サムネ生成の全コストが捨てられ続ける。「サムネ品質最優先」の施策が一切効かない'),
    (2, 'P0', 'GCP project 844705815004', 'OAuth同意画面を「テスト中」→「本番」へ公開',
     '公開の継続性', '09-20前後に稼働6chが一斉停止。09-09〜09-13の再来'),
    (3, 'P0', 'akashic-librarian ほか7ch', '再認可', '登録', '上位chが止まったまま'),
    (4, 'P1', 'ANTHROPIC_API_KEY', '有効なキーに差し替え', '台本品質',
     'dual生成が GPT単独に劣化。09-13夜の「有効化」は効いていない'),
]:
    for i, v in enumerate(row_, 1):
        cl = ws.cell(r, i, v); cl.border = B
        cl.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(r, 2).fill = RED if row_[1] == 'P0' else YEL
    ws.row_dimensions[r].height = 56; r += 1

# ============================ 前回比較 ============================
ws = wb.create_sheet('前回比較')
setw(ws, [26, 12, 12, 10, 12, 12, 10, 13, 13, 10, 16])
ws['A1'] = '前回比較'; ws['A1'].font = H1
ws['A2'] = ('比較軸: 09-08スナップショット → 09-13スナップショット。'
            '09-09〜09-12 は fetch が1件も存在しないため、日次の前日比は原理的に作れない')
ws['A2'].font = Font(italic=True, size=9, color='C00000')
r = 4
hrow(ws, r, ['チャンネル', '再生 09-08', '再生 09-13', '再生差', '登録 09-08', '登録 09-13',
             '登録差', '登録/千 09-08', '登録/千 09-13', '差', '判定']); r += 1
for ch in ['scp-lab', 'yokai-watch', 'company-facts', 'daily-science', '2ch-matome']:
    a = agg(rows, ch); b = agg(prow, ch)
    dp = a['spm'] - b['spm']
    verdict = '◎ 大きく改善' if dp >= 0.15 else (
        '○ 改善' if dp > 0.01 else ('× 悪化' if dp < -0.01 else '横ばい'))
    for i, v in enumerate([NAMES[ch], b['v'], a['v'], a['v'] - b['v'], b['sg'], a['sg'],
                           a['sg'] - b['sg'], round(b['spm'], 3), round(a['spm'], 3),
                           round(dp, 3), verdict], 1):
        ws.cell(r, i, v).border = B
    ws.cell(r, 10).fill = GRN if dp > 0.01 else (RED if dp < -0.01 else GRY)
    r += 1

r += 1
ws.cell(r, 1, '前回（09-12 / 09-13）施策の効果検証').font = Font(bold=True, size=12); r += 1
hrow(ws, r, ['実施日', '対象', '施策', '検証期日', '本日の判定', '次アクション']); r += 1
for row_ in [
    ('09-12', 'scp-lab', 'autopilot を平日限定→週7日（週15→21本）', '09-19',
     '△ 部分的に検証可。09-13以降3本公開され、本数の増加自体は稼働を確認。'
     'ただし全て views=0 で登録効果は未測定', '09-19に再判定。それまで本数を追加しない'),
    ('09-12', 'company-facts', 'autopilot枠 3→4（12:30増設・週21→28本）', '09-19',
     '△ 同上。09-13以降3本公開', '09-19に再判定'),
    ('09-12', '切り抜き3ch', 'ゲート付与＋キュー補充を title_constraints に通す', '—',
     '検証不能。3chとも停止中で公開0本', '再認可後'),
    ('09-13', '8ch', 'autopilot 無効化（5ch集中運用へ）', '09-20',
     '○ 成立。5chに絞った結果11本/2日を安定公開。生成リソースの分散は解消', '継続'),
    ('09-13', 'socio-rx', 'title_rules.hard_constraints を付与', '09-20',
     '検証不能。公開1本・再生0。ただし構造としてはゆっくり系5chが全てゲート下に入った', '初回実績待ち'),
    ('09-13', 'company-facts', '毎日投稿化（top-level days_of_week の地雷を解消）', '09-20',
     '○ 成立。09-13に2本＋23:15に1本。週3日制約は解けている', '継続'),
    ('09-13夜', '全ch', 'ゲート整備 全般', '—',
     '× 本日判明した重大な穴: ゲートは違反を検知していたのに、公開は止まらなかった。'
     '「ゲートを付けた」は「違反が止まる」を意味しない',
     '【A】-1 のコード修正で塞いだ。09-15に実効を確認'),
]:
    for i, v in enumerate(row_, 1):
        cl = ws.cell(r, i, v); cl.border = B
        cl.alignment = Alignment(wrap_text=True, vertical='top')
    j = row_[4][0]
    ws.cell(r, 5).fill = GRN if j == '○' else (
        YEL if j == '△' else (RED if j == '×' else GRY))
    ws.row_dimensions[r].height = 56; r += 1
ws.column_dimensions['C'].width = 40
ws.column_dimensions['E'].width = 54
ws.column_dimensions['F'].width = 34

r += 1
ws.cell(r, 1, 'データ鮮度の記録（同一スナップショットで2回configを変えない鉄則の根拠）'
        ).font = Font(bold=True, size=12); r += 1
hrow(ws, r, ['fetch日時(JST)', 'video_metrics行数', '対象ch数', 'このスナップショットを使った判断']); r += 1
for row_ in [
    ('2026-09-06 23:00', 362, 12, '09-06〜09-08 の分析'),
    ('2026-09-07 23:00', 340, 9, '—'),
    ('2026-09-08 23:00', 351, 9, '09-09〜09-13 の分析（5日間これしか無かった）'),
    ('2026-09-13 22:31〜22:40', 246, 6,
     '①09-13夜の指揮者（socio-rx / company-facts）'
     '②09-14朝のバックエンド自動更新（投稿時間・型優先度）③本日の本レポート'),
]:
    for i, v in enumerate(row_, 1):
        cl = ws.cell(r, i, v); cl.border = B
        cl.alignment = Alignment(wrap_text=True, vertical='top')
    r += 1
cl = ws.cell(r, 1, '→ 09-13 22:40 のスナップショットは既に2回 config 判断に使われている。'
                   'よって本日の指揮者は「実績にもとづく3回目の config 変更」を行わず、'
                   '検算とコミット、および実績に依存しない構造バグの修正に限定した。')
cl.font = Font(italic=True, size=9, color='C00000')
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)

wb.save(OUT)
print('saved', OUT, os.path.getsize(OUT))
