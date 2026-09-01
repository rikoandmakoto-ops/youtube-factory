#!/usr/bin/env python3
"""台本A/Bテスト集計 — Claude 台本 vs GPT 台本を実績で比較して xlsx に出す。

前提:
  - 出所の台帳は analytics.db の video_script_source（post_upload が投稿時に書く。
    過去分は backfill_script_source.py で埋める）
  - 実績は analytics.db の video_metrics（youtube_analytics.fetch_video_metrics）

比較する指標:
  - 視聴回数 (views) と 1日あたり視聴回数 (views/day) ※公開からの日数で割った値
  - 平均維持率 (avg_view_percentage)
  - いいね率 (likes / views)
  - チャンネル登録転換率 (subscribers_gained / views)
  - インプレッションCTR (ctr) ※ Reporting API 由来。取れていない動画は母数から除く

使い方:
  # 台帳にある動画の実績を API から取り直してから集計
  python ab_script_source_report.py --refresh

  # 保存済みスナップショットだけで集計（API を叩かない）
  python ab_script_source_report.py

  # 判定に必要な本数を変える（既定: 片側5本）
  python ab_script_source_report.py --min-n 5

出力:
  data/reports/ab_script_source_<YYYY-MM-DD>.xlsx
    - サマリ      … 全チャンネルの Claude vs GPT を1行ずつ
    - <channel>   … チャンネルごとの指標別比較表
    - 明細        … 集計に使った全動画の生データ

注意:
  - 本数が片側 min-n に満たないチャンネルは「サンプル不足」と出し、勝敗は書かない。
  - views は公開からの日数に強く依存するので、判定は views/day を見る。
  - CTR は Reporting API の取り込み待ち（YouTube 側で2〜3日遅れる）で 0 のことが
    ある。impressions=0 の動画は CTR の母数から外し、その本数を併記する。
"""

import argparse
import json
import os
import statistics
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

BACKEND_DIR = Path(__file__).parent
sys.path.insert(0, str(BACKEND_DIR))

env_file = BACKEND_DIR / ".env"
if env_file.exists():
    for line in env_file.read_text().splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        os.environ.setdefault(k.strip(), v.strip())

from pipeline.analytics import script_source as ss  # noqa: E402
from pipeline.analytics import store as analytics_store  # noqa: E402

PROJECT_ROOT = BACKEND_DIR.parent
REPORTS_DIR = PROJECT_ROOT / "data" / "reports"

CLAUDE = "claude"
GPT = "gpt"

# (キー, 表示名, 数値書式, 大きいほうが良いか)
METRICS: List[Tuple[str, str, str, bool]] = [
    ("views", "視聴回数", "#,##0.0", True),
    ("views_per_day", "1日あたり視聴回数", "#,##0.00", True),
    ("avg_view_percentage", "平均維持率", "0.0%", True),
    ("like_rate", "いいね率", "0.00%", True),
    ("subscribe_rate", "登録転換率", "0.000%", True),
    ("ctr", "インプレッションCTR", "0.00%", True),
]

# 判定に使う主指標（views は公開日数の影響が大きいので views/day を使う）
VERDICT_METRICS = ["views_per_day", "avg_view_percentage", "like_rate",
                   "subscribe_rate", "ctr"]

# 中央値の差がこの比率以下なら「差なし」とする
TIE_THRESHOLD = 0.05


def _pct(value: Optional[float]) -> Optional[float]:
    """0〜100 で来る率を 0〜1 に揃える（0〜1 のものはそのまま）。"""
    if value is None:
        return None
    v = float(value)
    return v / 100.0 if v > 1 else v


def _age_days(published_at: Optional[str]) -> Optional[float]:
    if not published_at:
        return None
    try:
        s = str(published_at).replace("Z", "+00:00")
        dt = datetime.fromisoformat(s)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
    except Exception:
        return None
    delta = datetime.now(timezone.utc) - dt
    return max(delta.total_seconds() / 86400.0, 0.0)


def _build_rows(*, refresh: bool, channels: Optional[List[str]]) -> List[Dict[str, Any]]:
    """台帳 × 実績スナップショットを突き合わせて、動画1本=1行にする。"""
    ledger = ss.list_for_channel(limit=10000)
    if channels:
        ledger = [r for r in ledger if r.get("channel_id") in channels]
    if not ledger:
        return []

    by_channel: Dict[str, List[Dict[str, Any]]] = {}
    for r in ledger:
        by_channel.setdefault(r["channel_id"], []).append(r)

    if refresh:
        from pipeline.youtube_analytics import fetch_video_metrics

        for ch, entries in by_channel.items():
            vids = [e["video_id"] for e in entries]
            print(f"  🔄 実績取得 [{ch}] {len(vids)}本 ...")
            try:
                res = fetch_video_metrics(ch, video_ids=vids, days=90)
                if not res.get("ok"):
                    print(f"     ⚠️ {res.get('error')}")
            except Exception as e:
                print(f"     ⚠️ 取得失敗: {e}")

    rows: List[Dict[str, Any]] = []
    for ch, entries in by_channel.items():
        metrics = {
            m["video_id"]: m
            for m in analytics_store.list_video_metrics(ch, limit=2000)
        }
        for e in entries:
            m = metrics.get(e["video_id"]) or {}
            views = float(m.get("views") or 0)
            published_at = e.get("published_at") or m.get("published_at")
            age = _age_days(published_at)
            impressions = int(m.get("impressions") or 0)
            rows.append(
                {
                    "channel_id": ch,
                    "video_id": e["video_id"],
                    "script_source": e["script_source"],
                    "title": e.get("title") or m.get("title") or "",
                    "published_at": published_at,
                    "age_days": round(age, 2) if age is not None else None,
                    "has_metrics": bool(m),
                    "views": views,
                    # 公開直後は 1日未満で割ると跳ねるので最低 1日として扱う
                    "views_per_day": (views / max(age, 1.0)) if age is not None else None,
                    "avg_view_percentage": _pct(m.get("avg_view_percentage")),
                    "like_rate": (float(m.get("likes") or 0) / views) if views else None,
                    "subscribe_rate": (
                        float(m.get("subscribers_gained") or 0) / views if views else None
                    ),
                    "impressions": impressions,
                    # impressions が 0 の動画は Reporting API 未取り込み。CTR は
                    # 「0%」ではなく「不明」なので母数から外す。
                    "ctr": _pct(m.get("ctr")) if impressions > 0 else None,
                    "likes": int(m.get("likes") or 0),
                    "subscribers_gained": int(m.get("subscribers_gained") or 0),
                    "snapshot_date": m.get("date"),
                    # views=0 のスナップショットは「本当に0再生」ではなく
                    # 「その計測窓に実績が入っていない」ことがほとんど（Analytics の
                    # 1〜3日遅れ、または古い日付のまま更新されていない行）。
                    # 0 として平均に混ぜると台帳の古い側だけが不当に下がるので、
                    # 集計母数から外して除外本数を別に出す。
                    "measured": bool(m) and views > 0,
                }
            )
    rows.sort(key=lambda r: (r["channel_id"], r.get("published_at") or ""))
    return rows


def _agg(values: List[float]) -> Dict[str, Any]:
    vals = [v for v in values if v is not None]
    if not vals:
        return {"n": 0, "median": None, "mean": None}
    return {
        "n": len(vals),
        "median": statistics.median(vals),
        "mean": statistics.fmean(vals),
    }


def _compare(channel_rows: List[Dict[str, Any]], min_n: int) -> Dict[str, Any]:
    """1チャンネル分の Claude vs GPT 比較を組む。

    集計・判定は「実績が計測できている本数」で行う。台帳の本数（ledger）は
    別に出す。両者が乖離しているときは、スナップショットの取り直し
    （--refresh）が足りていないというサインになる。
    """
    ledger = {
        src: [r for r in channel_rows if r["script_source"] == src]
        for src in (CLAUDE, GPT)
    }
    sides = {src: [r for r in rs if r["measured"]] for src, rs in ledger.items()}
    out: Dict[str, Any] = {
        "counts": {src: len(rs) for src, rs in sides.items()},
        "ledger_counts": {src: len(rs) for src, rs in ledger.items()},
        "unmeasured": {
            src: len(ledger[src]) - len(sides[src]) for src in (CLAUDE, GPT)
        },
        "metrics": {},
        "wins": {CLAUDE: 0, GPT: 0, "tie": 0},
    }

    for key, label, _fmt, higher_better in METRICS:
        agg = {src: _agg([r.get(key) for r in rs]) for src, rs in sides.items()}
        c_med, g_med = agg[CLAUDE]["median"], agg[GPT]["median"]
        winner = None
        diff_ratio = None
        enough = agg[CLAUDE]["n"] >= min_n and agg[GPT]["n"] >= min_n
        if c_med is not None and g_med is not None:
            base = max(abs(c_med), abs(g_med))
            diff_ratio = ((c_med - g_med) / base) if base else 0.0
            if not higher_better:
                diff_ratio = -diff_ratio
            if enough:
                if abs(diff_ratio) <= TIE_THRESHOLD:
                    winner = "tie"
                else:
                    winner = CLAUDE if diff_ratio > 0 else GPT
        out["metrics"][key] = {
            "label": label,
            "claude": agg[CLAUDE],
            "gpt": agg[GPT],
            "diff_ratio": diff_ratio,
            "winner": winner,
            "enough_samples": enough,
        }
        if winner and key in VERDICT_METRICS:
            out["wins"][winner] = out["wins"].get(winner, 0) + 1

    n_c, n_g = out["counts"][CLAUDE], out["counts"][GPT]
    if n_c < min_n or n_g < min_n:
        out["verdict"] = (
            f"サンプル不足 (計測済み claude {n_c}本 / gpt {n_g}本 — 片側{min_n}本必要)"
        )
    elif out["wins"][CLAUDE] > out["wins"][GPT]:
        out["verdict"] = f"Claude優勢 ({out['wins'][CLAUDE]}勝 vs {out['wins'][GPT]}勝)"
    elif out["wins"][GPT] > out["wins"][CLAUDE]:
        out["verdict"] = f"GPT優勢 ({out['wins'][GPT]}勝 vs {out['wins'][CLAUDE]}勝)"
    else:
        out["verdict"] = f"互角 ({out['wins'][CLAUDE]}勝 vs {out['wins'][GPT]}勝)"
    return out


def _write_xlsx(
    rows: List[Dict[str, Any]],
    per_channel: Dict[str, Dict[str, Any]],
    out_path: Path,
    min_n: int,
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    claude_fill = PatternFill("solid", fgColor="E8F0FE")
    gpt_fill = PatternFill("solid", fgColor="FDF0E8")

    def _style_header(ws, ncols: int, widths: Optional[List[int]] = None) -> None:
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            width = (widths[c - 1] if widths and c - 1 < len(widths) else None)
            ws.column_dimensions[get_column_letter(c)].width = width or max(
                14, len(str(cell.value or "")) + 2
            )
        ws.freeze_panes = "A2"

    wb = Workbook()

    # ── サマリ ──
    ws = wb.active
    ws.title = "サマリ"
    summary_cols = [
        "チャンネル", "Claude本数(計測済)", "GPT本数(計測済)",
        "Claude本数(台帳)", "GPT本数(台帳)", "判定",
        "1日あたり視聴回数(claude)", "1日あたり視聴回数(gpt)",
        "平均維持率(claude)", "平均維持率(gpt)",
        "いいね率(claude)", "いいね率(gpt)",
        "登録転換率(claude)", "登録転換率(gpt)",
        "CTR(claude)", "CTR(gpt)",
    ]
    ws.append(summary_cols)
    for ch in sorted(per_channel):
        cmp_ = per_channel[ch]
        m = cmp_["metrics"]

        def _med(key: str, src: str):
            return m[key][src]["median"]

        ws.append([
            ch,
            cmp_["counts"][CLAUDE], cmp_["counts"][GPT],
            cmp_["ledger_counts"][CLAUDE], cmp_["ledger_counts"][GPT],
            cmp_["verdict"],
            _med("views_per_day", CLAUDE), _med("views_per_day", GPT),
            _med("avg_view_percentage", CLAUDE), _med("avg_view_percentage", GPT),
            _med("like_rate", CLAUDE), _med("like_rate", GPT),
            _med("subscribe_rate", CLAUDE), _med("subscribe_rate", GPT),
            _med("ctr", CLAUDE), _med("ctr", GPT),
        ])
    for r in range(2, ws.max_row + 1):
        for c in (7, 8):
            ws.cell(row=r, column=c).number_format = "#,##0.00"
        for c in (9, 10):
            ws.cell(row=r, column=c).number_format = "0.0%"
        for c in (11, 12, 15, 16):
            ws.cell(row=r, column=c).number_format = "0.00%"
        for c in (13, 14):
            ws.cell(row=r, column=c).number_format = "0.000%"
    _style_header(ws, len(summary_cols), [20, 17, 15, 15, 13, 40] + [16] * 10)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(summary_cols))}{ws.max_row}"

    # 読み違いを防ぐための注記。特に過去分は「交互に出し分けた結果」ではない。
    notes = [
        "",
        ["読み方の注意"],
        ["・数値は中央値。views は公開からの日数に強く依存するので判定は"
         "1日あたり視聴回数で行う。"],
        [f"・片側 {min_n} 本未満のチャンネルは判定しない（サンプル不足と表示）。"],
        ["・CTR は Reporting API 由来。YouTube 側の集計が2〜3日遅れるため、"
         "直近の動画は impressions=0（不明）として母数から外している。"],
        ["・resolved_from=archive の過去分は『交互に出し分けたA/B』ではなく、"
         "生成時のブラインド評価で勝ったほうを採用した結果なので、"
         "モデル差とテーマ・時期の差が混ざっている。厳密な比較は"
         "台帳に交互投稿が溜まってからの分で行うこと。"],
        ["・2026-08-19〜25 に台本字数を 193字→372字→200字 と動かしているため、"
         "その期間をまたぐ比較には尺の影響も乗っている。"],
    ]
    for n in notes:
        ws.append(n if isinstance(n, list) else [n])
        if isinstance(n, list) and n and n[0] == "読み方の注意":
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)

    # ── チャンネル別 ──
    detail_cols = [
        "指標", "Claude中央値", "GPT中央値", "Claude平均", "GPT平均",
        "Claude本数", "GPT本数", "差(Claude基準)", "優勢",
    ]
    for ch in sorted(per_channel):
        cmp_ = per_channel[ch]
        ws = wb.create_sheet(title=ch[:31])
        ws.append(detail_cols)
        for key, label, fmt, _hb in METRICS:
            d = cmp_["metrics"][key]
            if d["winner"] == "tie":
                mark = "差なし"
            elif d["winner"] == CLAUDE:
                mark = "Claude"
            elif d["winner"] == GPT:
                mark = "GPT"
            else:
                mark = "判定不可"
            ws.append([
                label,
                d[CLAUDE]["median"], d[GPT]["median"],
                d[CLAUDE]["mean"], d[GPT]["mean"],
                d[CLAUDE]["n"], d[GPT]["n"],
                d["diff_ratio"], mark,
            ])
            r = ws.max_row
            for c in (2, 3, 4, 5):
                ws.cell(row=r, column=c).number_format = fmt
            ws.cell(row=r, column=8).number_format = "+0.0%;-0.0%;0.0%"
            ws.cell(row=r, column=2).fill = claude_fill
            ws.cell(row=r, column=4).fill = claude_fill
            ws.cell(row=r, column=3).fill = gpt_fill
            ws.cell(row=r, column=5).fill = gpt_fill
        ws.append([])
        ws.append(["判定", cmp_["verdict"]])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.append([
            "台帳本数",
            f"claude {cmp_['ledger_counts'][CLAUDE]}本 / gpt {cmp_['ledger_counts'][GPT]}本"
            f"（うち実績未計測: claude {cmp_['unmeasured'][CLAUDE]}本 / "
            f"gpt {cmp_['unmeasured'][GPT]}本 — --refresh で取得）",
        ])
        ws.append([
            "注記",
            f"判定は {', '.join(VERDICT_METRICS)} の中央値勝敗。"
            f"差が±{TIE_THRESHOLD:.0%}以内なら差なし。片側{min_n}本未満は判定しない。"
            "CTR は impressions が取れている動画のみが母数。"
            "views=0 のスナップショットは未計測とみなして母数から外している。",
        ])
        _style_header(ws, len(detail_cols), [20, 14, 14, 14, 14, 11, 10, 15, 10])

    # ── 明細 ──
    ws = wb.create_sheet(title="明細")
    raw_cols = [
        "チャンネル", "台本", "video_id", "タイトル", "公開日時", "経過日数",
        "視聴回数", "1日あたり視聴回数", "平均維持率", "いいね率", "登録転換率",
        "インプレッション", "CTR", "スナップショット日", "集計対象",
    ]
    ws.append(raw_cols)
    for r in rows:
        ws.append([
            r["channel_id"], r["script_source"], r["video_id"], r["title"],
            r["published_at"], r["age_days"], r["views"], r["views_per_day"],
            r["avg_view_percentage"], r["like_rate"], r["subscribe_rate"],
            r["impressions"], r["ctr"], r["snapshot_date"],
            "yes" if r["measured"] else ("no(0再生)" if r["has_metrics"] else "no(未取得)"),
        ])
        i = ws.max_row
        ws.cell(row=i, column=8).number_format = "#,##0.00"
        ws.cell(row=i, column=9).number_format = "0.0%"
        ws.cell(row=i, column=10).number_format = "0.00%"
        ws.cell(row=i, column=11).number_format = "0.000%"
        ws.cell(row=i, column=13).number_format = "0.00%"
        ws.cell(row=i, column=2).fill = claude_fill if r["script_source"] == CLAUDE else gpt_fill
    _style_header(ws, len(raw_cols), [20, 9, 14, 46, 22, 11, 11, 16, 12, 11, 12, 16, 10, 16, 13])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(raw_cols))}{ws.max_row}"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main() -> int:
    ap = argparse.ArgumentParser(description="Claude台本 vs GPT台本 の実績比較")
    ap.add_argument("--refresh", action="store_true",
                    help="集計前に YouTube Analytics から実績を取り直す")
    ap.add_argument("--min-n", type=int, default=5,
                    help="判定に必要な片側の本数（既定 5）")
    ap.add_argument("--channel", action="append",
                    help="対象チャンネル（複数可、既定は台帳にある全部）")
    ap.add_argument("--out", help="出力先 xlsx パス")
    ap.add_argument("--json", dest="json_out", help="集計結果を JSON でも書き出す")
    args = ap.parse_args()

    print("📊 台本A/Bテスト集計")
    rows = _build_rows(refresh=args.refresh, channels=args.channel)
    if not rows:
        print("台帳が空です。backfill_script_source.py で投稿分を記録してください。")
        return 1

    per_channel: Dict[str, Dict[str, Any]] = {}
    for ch in sorted({r["channel_id"] for r in rows}):
        per_channel[ch] = _compare([r for r in rows if r["channel_id"] == ch], args.min_n)

    print(f"\n{'チャンネル':<22}{'claude':>10}{'gpt':>10}  判定")
    print(f"{'':22}{'(計測/台帳)':>12}{'(計測/台帳)':>12}")
    ready = 0
    for ch, cmp_ in per_channel.items():
        c = f"{cmp_['counts'][CLAUDE]}/{cmp_['ledger_counts'][CLAUDE]}"
        g = f"{cmp_['counts'][GPT]}/{cmp_['ledger_counts'][GPT]}"
        print(f"{ch:<22}{c:>10}{g:>10}  {cmp_['verdict']}")
        if not cmp_["verdict"].startswith("サンプル不足"):
            ready += 1
    unmeasured = sum(1 for r in rows if not r["measured"])
    if unmeasured:
        no_row = sum(1 for r in rows if not r["has_metrics"])
        print(f"\n⚠️ 集計に使えない動画が {unmeasured} 本"
              f"（スナップショット未取得 {no_row} 本 / 0再生スナップショット "
              f"{unmeasured - no_row} 本）。--refresh で取り直せます。")

    date_str = datetime.now().strftime("%Y-%m-%d")
    out_path = Path(args.out) if args.out else REPORTS_DIR / f"ab_script_source_{date_str}.xlsx"
    _write_xlsx(rows, per_channel, out_path, args.min_n)
    print(f"\n📁 xlsx: {out_path}")
    print(f"   判定できたチャンネル: {ready}/{len(per_channel)}")

    if args.json_out:
        Path(args.json_out).write_text(
            json.dumps(
                {"generated_at": date_str, "min_n": args.min_n,
                 "per_channel": per_channel, "videos": rows},
                ensure_ascii=False, indent=2, default=str,
            ),
            encoding="utf-8",
        )
        print(f"   json: {args.json_out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
